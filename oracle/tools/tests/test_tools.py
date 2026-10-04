"""Tests for the Oracle metadata tooling. No network. Run with:

    python3 -m pytest oracle/tools/tests -q
"""

from __future__ import annotations

import json
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import build_data_dictionary as bdd  # noqa: E402
import osvc_client as oc  # noqa: E402
import pull_metadata as pm  # noqa: E402


# --- allowlist guard -------------------------------------------------------------

def test_schema_and_menu_paths_allowed():
    assert oc.check_path("metadata-catalog") == "metadata-catalog"
    assert oc.check_path("/metadata-catalog/incidents") == "metadata-catalog/incidents"
    assert oc.check_path("metadata-catalog/SCIF.SCIF") == "metadata-catalog/SCIF.SCIF"
    assert oc.check_path("namedIDs/incidents/severity") == "namedIDs/incidents/severity"
    assert oc.check_path("namedIDHierarchies/incidents/source") == "namedIDHierarchies/incidents/source"
    assert oc.check_path("serviceCategories") == "serviceCategories"
    assert oc.check_path("accountGroups") == "accountGroups"
    assert oc.check_path("analyticsReports") == "analyticsReports"
    for added in ("eventSubscriptions", "mailboxes", "serviceMailboxes", "standardContents"):
        assert oc.check_path(added) == added


def test_record_query_and_report_run_resources_refused():
    for bad in (
        "incidents", "incidents/12", "incidentResponse", "contacts", "tasks",
        "accounts", "chats", "queryResults", "analyticsReportResults",
        "configurations", "bulkExtracts",
    ):
        try:
            oc.check_path(bad)
        except oc.RefusedRequest as exc:
            assert "Refused" in str(exc)
        else:
            raise AssertionError(f"{bad} was allowed")


def test_custom_object_rows_need_run_time_approval():
    try:
        oc.check_path("SCIF.mGender")
    except oc.RefusedRequest as exc:
        assert "not classified as a menu-only table" in str(exc)
    else:
        raise AssertionError("unapproved custom object was allowed")
    assert oc.check_path("SCIF.mGender", approved_menu_objects={"SCIF.mGender"}) == "SCIF.mGender"
    try:
        oc.check_path("SCIF.SCIF", approved_menu_objects={"SCIF.mGender"})
    except oc.RefusedRequest:
        pass
    else:
        raise AssertionError("SCIF.SCIF was allowed")


def test_unknown_and_malformed_paths_refused():
    for bad in ("", "../metadata-catalog", "metadata-catalog?query=x", "somethingNew", "messageBases"):
        try:
            oc.check_path(bad)
        except oc.RefusedRequest:
            pass
        else:
            raise AssertionError(f"{bad!r} was allowed")


def test_client_refuses_before_any_network_call():
    class ExplodingSession:
        headers: dict = {}

        def get(self, *a, **k):
            raise AssertionError("network call attempted")

    with tempfile.TemporaryDirectory() as tmp:
        client = oc.OsvcClient("https://example.invalid", ("u", "p"), Path(tmp) / "m.jsonl", session=ExplodingSession())
        for bad in ("incidents", "SCIF.SCIF", "analyticsReportResults"):
            try:
                client.get(bad)
            except oc.RefusedRequest:
                pass
            else:
                raise AssertionError(f"{bad} reached the network")
        try:
            client.get("serviceCategories", params={"q": "x"})
        except oc.RefusedRequest:
            pass
        else:
            raise AssertionError("arbitrary query parameter was allowed")
        assert not (Path(tmp) / "m.jsonl").exists()


def test_answer_records_allowed_because_articles_migrate():
    # Ben Bolding, 2026-10-01: knowledge base articles are in migration scope and hold no caller PII.
    for ok in ("answers", "answers/12", "answerVersions"):
        assert oc.check_path(ok) == ok


def test_answer_query_guard_allows_answer_queries_only():
    for ok in (
        "SELECT COUNT(*) FROM Answer",
        "SELECT AnswerType.LookupName, CustomFields.c.answer_types.LookupName, COUNT(*) FROM Answer "
        "GROUP BY AnswerType.LookupName, CustomFields.c.answer_types.LookupName",
        "select A.ID from Answer A where A.Language.LookupName = 'es_ES'",
        "SELECT answerType.lookupName, COUNT(*) FROM answers GROUP BY answerType.lookupName",
    ):
        assert oc.check_answer_query(ok) == ok.strip()
    for bad in (
        "",
        "SELECT * FROM Incident",
        "SELECT Name.First FROM Contact",
        "SELECT ID FROM Answer; SELECT * FROM Contact",
        "SELECT UpdatedByAccount.Name FROM Answer",
        "SELECT ID FROM Answer WHERE ID IN (SELECT Answer FROM Incident)",
        "SELECT A.ID FROM Answer A, Chat C",
    ):
        try:
            oc.check_answer_query(bad)
        except oc.RefusedRequest as exc:
            assert "Refused" in str(exc)
        else:
            raise AssertionError(f"{bad!r} was allowed")


def test_query_answers_sends_query_and_logs():
    calls = []

    class FakeSession:
        headers: dict = {}

        def get(self, url, auth=None, timeout=None, headers=None, params=None):
            calls.append((url, params))
            return _FakeResp({"items": [{"tableName": "Answer", "columnNames": ["count(*)"], "rows": [["5"]]}]})

    with tempfile.TemporaryDirectory() as tmp:
        manifest = Path(tmp) / "m.jsonl"
        client = oc.OsvcClient("https://example.invalid", ("u", "p"), manifest, session=FakeSession())
        status, body = client.query_answers("SELECT COUNT(*) FROM Answer")
        assert status == 200 and body["items"][0]["rows"] == [["5"]]
        assert calls[0][0].endswith("/queryResults") and calls[0][1] == {"query": "SELECT COUNT(*) FROM Answer"}
        assert len(manifest.read_text().splitlines()) == 1
        try:
            client.query_answers("SELECT * FROM Contact")
        except oc.RefusedRequest:
            pass
        else:
            raise AssertionError("contact query was sent")
        assert len(calls) == 1


class _FakeResp:
    def __init__(self, body: dict, status: int = 200):
        self.status_code = status
        self.content = json.dumps(body).encode()
        self.headers = {"Content-Type": "application/json"}
        self.url = "https://example.invalid/x"


def test_manifest_and_accept_header_on_allowed_call():
    calls = []

    class FakeSession:
        headers: dict = {}

        def get(self, url, auth=None, timeout=None, headers=None, params=None):
            calls.append((url, headers, params))
            return _FakeResp({"items": []})

    with tempfile.TemporaryDirectory() as tmp:
        manifest = Path(tmp) / "m.jsonl"
        client = oc.OsvcClient("https://example.invalid", ("u", "p"), manifest, session=FakeSession())
        status, body = client.get_json("metadata-catalog")
        assert status == 200 and body == {"items": []}
        assert calls[0][1] == {"Accept": "application/json"}
        client.get_json("metadata-catalog/incidents", accept=oc.SCHEMA_JSON)
        assert calls[1][1] == {"Accept": "application/schema+json"}
        entries = [json.loads(line) for line in manifest.read_text().splitlines()]
        assert len(entries) == 2 and entries[0]["status"] == 200 and "sha256" in entries[0]


def test_pagination_follows_has_more():
    pages = [
        {"items": [{"id": 1, "lookupName": "A"}], "hasMore": True},
        {"items": [{"id": 2, "lookupName": "B"}], "hasMore": False},
    ]
    seen_offsets = []

    class FakeSession:
        headers: dict = {}

        def get(self, url, auth=None, timeout=None, headers=None, params=None):
            seen_offsets.append(params["offset"])
            return _FakeResp(pages[len(seen_offsets) - 1])

    with tempfile.TemporaryDirectory() as tmp:
        client = oc.OsvcClient("https://example.invalid", ("u", "p"), Path(tmp) / "m.jsonl", session=FakeSession())
        client.approve_menu_object("SCIF.mGender")
        status, items, raw = client.get_all_items("SCIF.mGender", page_size=1)
        assert status == 200 and [i["id"] for i in items] == [1, 2]
        assert seen_offsets == [0, 1] and len(raw) == 2


# --- classification and exclusions ---------------------------------------------------

MENU_SCHEMA = {"properties": {"id": {"type": "integer"}, "lookupName": {"type": "string"}, "displayOrder": {"type": "integer"}, "createdTime": {}, "updatedTime": {}, "links": {}}}
DATA_SCHEMA = {"properties": {"id": {}, "lookupName": {}, "quitDate": {"type": "string", "format": "date"}, "createdTime": {}}}
ALLOF_MENU_SCHEMA = {"allOf": [{"properties": {"id": {}, "lookupName": {}}}, {"properties": {"names": {}}}]}


def test_classify_custom_object():
    assert oc.classify_custom_object(MENU_SCHEMA) == "menu"
    assert oc.classify_custom_object(DATA_SCHEMA) == "data"
    assert oc.classify_custom_object(ALLOF_MENU_SCHEMA) == "menu"
    assert oc.classify_custom_object({"$schema": "x", "allOf": [{"$ref": "y"}]}) == "unknown"
    assert oc.classify_custom_object(None) == "unknown"


def test_people_fields_excluded():
    for f in ("assignedTo", "createdByAccount", "updatedByAccount", "primaryContact", "organization", "owner"):
        assert oc.is_people_field(f), f
    for f in ("severity", "statusWithType", "source", "queue", "language", "mailbox", "channel", "contactType", "staffGroup", "profile"):
        assert not oc.is_people_field(f), f


def test_named_id_listing_parser():
    listing = {
        "items": [
            {"name": "severity", "links": [{"rel": "canonical", "href": "https://x/services/rest/connect/v1.4/namedIDs/incidents/severity"}]},
            {"name": "assignedTo", "links": [{"rel": "canonical", "href": "https://x/services/rest/connect/v1.4/namedIDs/incidents/assignedTo"}]},
        ],
        "links": [{"rel": "self", "href": "https://x/services/rest/connect/v1.4/namedIDs/incidents"}],
    }
    assert pm.field_names_from_named_id_listing(listing, "incidents") == ["severity", "assignedTo"]
    links_only = {"links": [{"rel": "x", "href": "https://x/services/rest/connect/v1.4/namedIDs/incidents/queue"}]}
    assert pm.field_names_from_named_id_listing(links_only, "incidents") == ["queue"]


def test_resource_names_from_catalog_shapes():
    listing = {"items": [{"name": "incidents", "links": [{"rel": "canonical", "href": "https://x/metadata-catalog/incidents"}]}, {"name": "SCIF.mGender"}]}
    assert pm.resource_names_from_catalog(listing) == ["incidents", "SCIF.mGender"]


def test_schema_scope_counts_match_plan():
    fixed = [r for group in pm.SCHEMA_SCOPE.values() for r in group]
    assert len(fixed) == 26 and len(set(fixed)) == 26  # 11 core, 9 menus, 6 site configuration; plus 21 custom = 47


# --- data dictionary and picklists ---------------------------------------------------

INCIDENT_SCHEMA = {
    "type": "object",
    "properties": {
        "id": {"type": "integer", "readOnly": True, "description": "ID"},
        "subject": {"type": "string", "maxLength": 240},
        "createdTime": {"type": "string", "format": "date-time", "readOnly": True},
        "statusWithType": {"type": "object", "properties": {"status": {"$ref": "https://x/metadata-catalog/namedIDs"}, "statusType": {"$ref": "https://x/metadata-catalog/namedIDs"}}},
        "primaryContact": {"$ref": "https://x/metadata-catalog/contacts"},
        "customFields": {"type": "object", "properties": {
            "c": {"type": "object", "properties": {"point_of_access": {"type": "string", "maxLength": 40}, "sms_consent": {"type": "boolean"}, "ecrf_notes": {"type": "string"}}},
            "CO": {"type": "object", "properties": {"intake_ref": {"type": "integer"}}},
        }},
    },
}


def test_custom_fields_are_walked_and_classified():
    by_path = {r["field_path"]: r for r in bdd.rows_for_schema("incidents", INCIDENT_SCHEMA)}
    assert by_path["customFields.c.point_of_access"]["classification"] == "custom (c)"
    assert by_path["customFields.c.point_of_access"]["suggested_sf_type"] == "Text(40)"
    assert by_path["customFields.c.sms_consent"]["suggested_sf_type"] == "Checkbox"
    assert by_path["customFields.CO.intake_ref"]["classification"] == "custom attribute (CO)"
    assert "customFields.c" not in by_path and "customFields" not in by_path


def test_menu_source_and_lookups():
    by_path = {r["field_path"]: r for r in bdd.rows_for_schema("incidents", INCIDENT_SCHEMA)}
    assert by_path["statusWithType.status"]["suggested_sf_type"] == "Picklist"
    assert by_path["primaryContact"]["suggested_sf_type"] == "Lookup"
    assert all(r["menu_source"] == "" for r in by_path.values())  # only build() sets it, from saved files


def test_build_outputs_dictionary_and_picklists():
    with tempfile.TemporaryDirectory() as tmp:
        extract = Path(tmp) / "2026-09-23"
        (extract / "schema").mkdir(parents=True)
        (extract / "schema" / "incidents.json").write_text(json.dumps(INCIDENT_SCHEMA))
        (extract / "named-ids").mkdir()
        (extract / "named-ids" / "incidents.severity.json").write_text(json.dumps({"items": [{"id": 1, "lookupName": "High"}, {"id": 2, "lookupName": "Low"}]}))
        (extract / "named-ids" / "hierarchy.incidents.source.json").write_text(json.dumps({"items": [{"id": 3, "lookupName": "Phone", "parents": [{"id": 9, "lookupName": "Root"}]}]}))
        (extract / "named-ids" / "_listing.incidents.json").write_text("{}")
        (extract / "menu-objects").mkdir()
        (extract / "menu-objects" / "SCIF.mGender.json").write_text(json.dumps({"items": [{"id": 1, "lookupName": "Female", "displayOrder": 1}], "pages": 1}))
        (extract / "menus").mkdir()
        (extract / "menus" / "analyticsReports.json").write_text(json.dumps({"items": [{"id": 5, "lookupName": "Callback report"}]}))
        (extract / "_scope.json").write_text(json.dumps({"custom_object_classification": {"SCIF.mGender": "menu"}}))

        df = bdd.build(extract)
        assert df.columns == bdd.COLUMNS and df.filter(df["object"] == "incidents").height == 10

        pk = bdd.build_picklists(extract)
        assert pk.columns == bdd.PICKLIST_COLUMNS
        assert pk.height == 4  # 2 severity + 1 hierarchy + 1 menu object; report definitions excluded
        assert set(pk["source"].to_list()) == {"namedID", "namedIDHierarchy", "menuObject"}
        assert pk.filter(pk["source"] == "namedIDHierarchy")["parent"].to_list() == ["Root"]

        bdd.main(["x", str(extract)])
        for name in ("data-dictionary.parquet", "data-dictionary.csv", "picklist-values.parquet", "picklist-values.csv"):
            assert (extract / name).exists(), name


def _iface_links(resource: str, row_id: int, iface_ids: list[int]) -> dict:
    base = f"https://x.example/services/rest/connect/v1.4/{resource}/{row_id}/adminVisibleInterfaces"
    return {"items": [{"rel": "canonical", "href": f"{base}/{i}"} for i in iface_ids]}


def test_service_menu_hierarchy_from_full_rows():
    """Full product/category/disposition rows replace the flat listing and carry the tree."""
    with tempfile.TemporaryDirectory() as tmp:
        extract = Path(tmp) / "2026-09-23"
        (extract / "menus").mkdir(parents=True)
        (extract / "menus" / "serviceDispositions.json").write_text(json.dumps({"items": [
            {"id": 9, "lookupName": "Break Off"}, {"id": 11, "lookupName": "Time Constraint"}, {"id": 412, "lookupName": "Spanish Transfer"}]}))
        (extract / "menus" / "siteInterfaces.json").write_text(json.dumps({"items": [
            {"id": 1, "lookupName": "nci"}, {"id": 2, "lookupName": "nci1"}, {"id": 5, "lookupName": "nci2"}]}))
        (extract / "config-rows").mkdir()
        (extract / "config-rows" / "serviceDispositions.json").write_text(json.dumps([
            {"id": 9, "lookupName": "Break Off", "displayOrder": 2, "parent": None,
             "adminVisibleInterfaces": _iface_links("serviceDispositions", 9, [1, 2, 5])},
            {"id": 11, "lookupName": "Time Constraint", "displayOrder": 1, "parent": {"id": 9, "lookupName": "Break Off"},
             "adminVisibleInterfaces": _iface_links("serviceDispositions", 11, [1, 2, 5])},
            {"id": 412, "lookupName": "Spanish Transfer", "displayOrder": 3, "parent": {"id": 11, "lookupName": "Time Constraint"},
             "adminVisibleInterfaces": _iface_links("serviceDispositions", 412, [1, 5])},
        ]))

        pk = bdd.build_picklists(extract)
        assert pk.columns == bdd.PICKLIST_COLUMNS
        disp = {r["lookup_name"]: r for r in pk.filter(pk["object"] == "serviceDispositions").iter_rows(named=True)}
        assert len(disp) == 3  # full rows replace the listing, no duplicates
        assert disp["Break Off"]["parent"] == "" and disp["Break Off"]["hierarchy_path"] == "Break Off"
        assert disp["Time Constraint"]["parent"] == "Break Off"
        assert disp["Spanish Transfer"]["hierarchy_path"] == "Break Off > Time Constraint > Spanish Transfer"
        assert disp["Spanish Transfer"]["visible_interfaces"] == "nci, nci2"
        assert disp["Time Constraint"]["display_order"] == 1
        # a menu with no full rows keeps its listing and leaves the new columns empty
        iface = pk.filter(pk["object"] == "siteInterfaces")
        assert iface.height == 3 and set(iface["hierarchy_path"].to_list()) == {""}


# --- end-to-end dry run of the scoped pull against a fake server ------------------------

REAL_CATALOG = Path(__file__).resolve().parents[2] / "extracts" / "oracle-metadata" / "2026-09-23" / "catalog.json"


def _fake_server(url: str, headers: dict | None, params: dict | None, calls: list) -> _FakeResp:
    """Replays the real catalog listing and canned responses for everything else."""
    rel = url.split("/services/rest/connect/v1.4/", 1)[1]
    calls.append(rel)
    accept = (headers or {}).get("Accept")
    if rel == "metadata-catalog":
        return _FakeResp(json.loads(REAL_CATALOG.read_text()))
    if rel.startswith("metadata-catalog/"):
        assert accept == "application/schema+json", rel
        name = rel.split("/", 1)[1]
        if name.startswith("SCIF.m") or name.startswith("Referrals.m"):
            return _FakeResp(MENU_SCHEMA)
        if "." in name:
            return _FakeResp(DATA_SCHEMA)
        return _FakeResp(INCIDENT_SCHEMA if name == "incidents" else {"properties": {"id": {}, "lookupName": {}, "name": {}}})
    if rel.startswith("namedIDs/") and rel.count("/") == 1:
        res = rel.split("/")[1]
        return _FakeResp({"items": [
            {"name": "severity", "links": [{"rel": "canonical", "href": f"{url}/severity"}]},
            {"name": "assignedTo", "links": [{"rel": "canonical", "href": f"{url}/assignedTo"}]},
        ]})
    if rel.startswith("namedIDs/"):
        assert not rel.endswith("/assignedTo"), "people field was fetched"
        return _FakeResp({"items": [{"id": 1, "lookupName": "High"}]})
    if rel.startswith("namedIDHierarchies/") and rel.count("/") == 1:
        return _FakeResp({"links": [{"rel": "x", "href": f"{url}/source"}]})
    if rel.startswith("namedIDHierarchies/"):
        return _FakeResp({"items": [{"id": 3, "lookupName": "Phone", "parents": []}]})
    if rel == "analyticsReports":
        return _FakeResp({"error": "forbidden"}, status=403)
    # menu structure resources, accountGroups, approved menu objects
    return _FakeResp({"items": [{"id": 1, "lookupName": "Value"}], "hasMore": False})


def test_scoped_pull_end_to_end_against_fake_server(monkeypatch):
    assert REAL_CATALOG.exists(), "needs the real catalog.json from the 2026-09-23 connection test"
    calls: list[str] = []

    class FakeSession:
        headers: dict = {}

        def get(self, url, auth=None, timeout=None, headers=None, params=None):
            return _fake_server(url, headers, params, calls)

    with tempfile.TemporaryDirectory() as tmp:
        monkeypatch.setattr(pm, "EXTRACTS_DIR", Path(tmp))
        monkeypatch.setattr(pm, "load_site_url", lambda: "https://example.invalid")
        monkeypatch.setattr(pm, "load_credentials", lambda: ("u", "p"))
        real_client = pm.OsvcClient
        monkeypatch.setattr(pm, "OsvcClient", lambda site, auth, manifest_path: real_client(site, auth, manifest_path, session=FakeSession()))

        assert pm.run(test_only=False) == 0

        out = next(Path(tmp).glob("*"))
        scope = json.loads((out / "_scope.json").read_text())
        assert len(scope["schemas_requested"]) == 47
        assert len(scope["schemas_skipped"]) == 16
        assert scope["schemas_not_listed"] == []
        kinds = scope["custom_object_classification"]
        assert len(kinds) == 21
        assert kinds["SCIF.mGender"] == "menu" and kinds["SCIF.SCIF"] == "data" and kinds["DEMOGR.Study"] == "data"
        assert scope["named_id_fields_excluded_people"]["incidents"] == ["assignedTo"]
        assert scope["analyticsReports_status"] == 403

        # files land where the dictionary builder expects them
        assert len(list((out / "schema").glob("*.json"))) == 47
        menu_files = {p.stem for p in (out / "menu-objects").glob("*.json")}
        assert menu_files == {k for k, v in kinds.items() if v == "menu"}
        assert (out / "named-ids" / "incidents.severity.json").exists()
        assert (out / "named-ids" / "hierarchy.incidents.source.json").exists()
        assert (out / "menus" / "accountGroups.json").exists()

        # nothing refused was ever requested, and no data-classified custom object rows either
        refused_heads = set(oc.REFUSED_PREFIXES) | {k for k, v in kinds.items() if v != "menu"}
        for rel in calls:
            head = rel.split("/", 1)[0]
            assert head not in refused_heads, rel
        manifest = [json.loads(l) for l in (out / "_manifest.jsonl").read_text().splitlines()]
        assert len(manifest) == len(calls)

        # the dictionary and picklist build runs on the dry-run output
        bdd.main(["x", str(out)])
        assert (out / "data-dictionary.csv").exists() and (out / "picklist-values.csv").exists()



# --- Oracle's real schema shape, menu guards, custom fields -------------------------------

def _oracle_wrapped(props: dict, is_menu: bool | None = None) -> dict:
    """The shape Oracle returns: a collection schema with the record under definitions.singularResource."""
    sr = {"type": "object", "allOf": [{"$ref": "x"}], "properties": props}
    if is_menu is not None:
        sr["isMenu"] = is_menu
    return {"$schema": "x", "type": "object", "properties": {"items": {"type": "array"}}, "definitions": {"singularResource": sr}}


ORACLE_MENU_OBJECT = _oracle_wrapped({"id": {}, "lookupName": {}, "DisplayOrder": {}, "Labels": {}, "createdTime": {}, "updatedTime": {}}, is_menu=True)
ORACLE_DATA_OBJECT = _oracle_wrapped({"id": {}, "lookupName": {}, "QuitDate": {"type": ["string", "null"], "format": "date"}, "Gender": {"$ref": "https://x/metadata-catalog/SCIF.mGender", "isEnumerable": True}, "CreatedByAccount": {}}, is_menu=False)


def test_classifier_reads_oracle_wrapped_schemas():
    assert oc.classify_custom_object(ORACLE_MENU_OBJECT) == "menu"
    assert oc.classify_custom_object(ORACLE_DATA_OBJECT) == "data"
    # isMenu false wins even when the fields look like a menu
    assert oc.classify_custom_object(_oracle_wrapped({"id": {}, "lookupName": {}}, is_menu=False)) == "data"
    # a menu flag never overrides real data fields
    assert oc.classify_custom_object(_oracle_wrapped({"id": {}, "lookupName": {}, "Notes": {}}, is_menu=True)) == "data"


def test_non_menu_keys_guard():
    assert oc.non_menu_keys([{"id": 1, "lookupName": "A", "links": []}]) == set()
    assert oc.non_menu_keys([{"id": 1, "lookupName": "A", "parents": [], "DisplayOrder": 1}]) == set()
    assert oc.non_menu_keys([{"id": 1, "lookupName": "A", "emails": ["x"]}]) == {"emails"}
    assert oc.non_menu_keys("nope") == {"<not a list>"}


def test_save_menu_rejects_non_menu_responses():
    with tempfile.TemporaryDirectory() as tmp:
        errors: dict = {}
        ok = pm.save_menu(Path(tmp) / "a.json", {"items": [{"id": 1, "lookupName": "A"}]}, errors, "k1")
        assert ok == 1 and (Path(tmp) / "a.json").exists()
        bad = pm.save_menu(Path(tmp) / "b.json", {"items": [{"id": 1, "lookupName": "A", "subject": "caller text"}]}, errors, "k2")
        assert bad == 0 and not (Path(tmp) / "b.json").exists()
        assert "subject" in errors["k2"] and "caller text" not in errors["k2"]


def test_dictionary_reads_wrapped_schema_and_custom_field_schemas():
    with tempfile.TemporaryDirectory() as tmp:
        extract = Path(tmp) / "d"
        (extract / "schema").mkdir(parents=True)
        (extract / "schema" / "incidents.json").write_text(json.dumps(_oracle_wrapped({
            "id": {"type": "integer", "isAvailableForPATCH": False, "label": "ID"},
            "severity": {"type": ["object", "null"], "$ref": "https://x/metadata-catalog/namedIDs", "isEnumerable": True, "label": "Severity"},
            "customFields": {"type": "object", "$ref": "https://x/metadata-catalog/incidents/customFields"},
        })))
        (extract / "schema" / "SCIF.SCIF.json").write_text(json.dumps(ORACLE_DATA_OBJECT))
        (extract / "schema-sub").mkdir()
        (extract / "schema-sub" / "incidents.customFields.json").write_text(json.dumps({"properties": {"c": {}}}))
        (extract / "schema-sub" / "incidents.customFields.c.json").write_text(json.dumps({"properties": {
            "special_code_1": {"type": ["object", "null"], "isEnumerable": True, "label": "Special Code 1", "$ref": "https://x/metadata-catalog/incidents/customFields/c/special_code_1"},
            "notes_text": {"type": ["string", "null"], "maxLength": 4000, "label": "Notes"},
        }}))
        (extract / "named-ids").mkdir()
        (extract / "named-ids" / "incidents.severity.json").write_text(json.dumps({"items": [{"id": 1, "lookupName": "High"}]}))
        (extract / "named-ids" / "incidents.customFields.c.special_code_1.json").write_text(json.dumps({"items": [{"id": 7, "lookupName": "Code A"}, {"id": 8, "lookupName": "Code B"}]}))
        (extract / "menu-objects").mkdir()
        (extract / "menu-objects" / "SCIF.mGender.json").write_text(json.dumps({"items": [{"id": 1, "lookupName": "Female"}]}))

        df = bdd.build(extract)
        rows = {(r["object"], r["field_path"]): r for r in df.to_dicts()}
        assert ("incidents", "customFields") not in rows
        assert rows[("incidents", "id")]["read_only"] is True
        sev = rows[("incidents", "severity")]
        assert sev["is_menu"] and sev["suggested_sf_type"] == "Picklist" and sev["menu_source"] == "named-ids/incidents.severity"
        sc = rows[("incidents", "customFields.c.special_code_1")]
        assert sc["classification"] == "custom (c)" and sc["label"] == "Special Code 1"
        assert sc["menu_source"] == "named-ids/incidents.customFields.c.special_code_1"
        assert rows[("incidents", "customFields.c.notes_text")]["suggested_sf_type"] == "Text Area (Long)"
        g = rows[("SCIF.SCIF", "Gender")]
        assert g["menu_source"] == "menu-objects/SCIF.mGender" and g["suggested_sf_type"] == "Picklist"
        assert rows[("SCIF.SCIF", "QuitDate")]["suggested_sf_type"] == "Date"

        pk = bdd.build_picklists(extract)
        custom = pk.filter(pk["source"] == "customMenu")
        assert custom.height == 2 and set(custom["field_or_object"]) == {"customFields.c.special_code_1"}


def test_topup_end_to_end_against_fake_server(monkeypatch):
    """Top-up on a folder from an earlier pull: menu objects, custom field schemas and
    options, nested menus. Checks nothing forbidden is requested and relationships are skipped."""
    calls: list[str] = []

    def server(url, headers, params):
        rel = url.split("/services/rest/connect/v1.4/", 1)[1]
        calls.append(rel)
        if rel.endswith("/customFields") and rel.startswith("metadata-catalog/"):
            return _FakeResp({"properties": {"c": {}, "CO": {}}})
        if rel == "metadata-catalog/incidents/customFields/c":
            return _FakeResp({"properties": {
                "special_code_1": {"isEnumerable": True, "$ref": "https://x/services/rest/connect/v1.4/metadata-catalog/incidents/customFields/c/special_code_1"},
                "free_text": {"type": ["string", "null"]},
            }})
        if rel == "metadata-catalog/incidents/customFields/CO":
            return _FakeResp({"properties": {
                "scif_link": {"isEnumerable": True, "$ref": "https://x/services/rest/connect/v1.4/metadata-catalog/SCIF.SCIF"},
            }})
        if rel.startswith("metadata-catalog/"):
            return _FakeResp({"properties": {}})
        if rel == "namedIDs/incidents/customFields/c/special_code_1":
            return _FakeResp({"items": [{"id": 1, "lookupName": "A"}, {"id": 2, "lookupName": "B"}]})
        if rel in ("namedIDs/incidents/statusWithType/status", "namedIDs/incidents/statusWithType/statusType"):
            return _FakeResp({"items": [{"id": 1, "lookupName": "Solved"}]})
        if rel == "namedIDs/contacts/contactType":
            return _FakeResp({"items": [{"id": 1, "lookupName": "Caller"}]})
        if rel == "SCIF.mGender":
            return _FakeResp({"items": [{"id": 1, "lookupName": "Female"}], "hasMore": False})
        raise AssertionError(f"unexpected request {rel}")

    class FakeSession:
        headers: dict = {}

        def get(self, url, auth=None, timeout=None, headers=None, params=None):
            return server(url, headers, params)

    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp) / "2026-09-23"
        (out / "schema").mkdir(parents=True)
        (out / "schema" / "SCIF.mGender.json").write_text(json.dumps(ORACLE_MENU_OBJECT))
        (out / "schema" / "SCIF.SCIF.json").write_text(json.dumps(ORACLE_DATA_OBJECT))
        (out / "named-ids").mkdir()
        (out / "named-ids" / "incidents.statusWithType.json").write_text(json.dumps({"status": {"links": []}, "statusType": {"links": []}, "links": []}))
        (out / "_scope.json").write_text("{}")

        monkeypatch.setattr(pm, "EXTRACTS_DIR", Path(tmp))
        monkeypatch.setattr(pm, "load_site_url", lambda: "https://example.invalid")
        monkeypatch.setattr(pm, "load_credentials", lambda: ("u", "p"))
        real_client = pm.OsvcClient
        monkeypatch.setattr(pm, "OsvcClient", lambda site, auth, manifest_path: real_client(site, auth, manifest_path, session=FakeSession()))

        assert pm.topup("2026-09-23") == 0

        scope = json.loads((out / "_scope.json").read_text())
        assert scope["custom_object_classification"] == {"SCIF.mGender": "menu", "SCIF.SCIF": "data"}
        assert (out / "menu-objects" / "SCIF.mGender.json").exists()
        assert not (out / "menu-objects" / "SCIF.SCIF.json").exists()
        assert (out / "named-ids" / "incidents.customFields.c.special_code_1.json").exists()
        assert scope["custom_enumerable_fields_skipped_as_relationships"] == ["incidents.customFields.CO.scif_link"]
        assert (out / "named-ids" / "incidents.statusWithType.status.json").exists()
        assert (out / "named-ids" / "contacts.contactType.json").exists()
        for rel in calls:
            head = rel.split("/", 1)[0]
            assert head not in oc.REFUSED_PREFIXES and head != "SCIF.SCIF", rel
        assert "namedIDs/incidents/customFields/CO/scif_link" not in calls
        assert not any("free_text" in c for c in calls)


def test_name_field_counts_as_menu_only_when_oracle_flags_it():
    named = {"id": {}, "lookupName": {}, "Name": {}, "DisplayOrder": {}, "createdTime": {}, "updatedTime": {}}
    assert oc.classify_custom_object(_oracle_wrapped(named, is_menu=True)) == "menu"
    assert oc.classify_custom_object(_oracle_wrapped(named)) == "data"
    assert oc.classify_custom_object(_oracle_wrapped(named, is_menu=False)) == "data"


def test_booleans_staff_lookups_and_menu_resources():
    with tempfile.TemporaryDirectory() as tmp:
        extract = Path(tmp) / "d"
        (extract / "schema").mkdir(parents=True)
        (extract / "schema" / "incidents.json").write_text(json.dumps(_oracle_wrapped({
            "AfterMeals": {"type": ["boolean", "null"], "isEnumerable": True},
            "createdByAccount": {"type": ["object", "null"], "isEnumerable": True, "$ref": "https://x/metadata-catalog/accounts"},
            "category": {"type": ["object", "null"], "isEnumerable": True, "$ref": "https://x/metadata-catalog/serviceCategories"},
            "source": {"type": "object", "$ref": "https://x/metadata-catalog/source"},
            "statusWithType": {"type": "object", "$ref": "https://x/metadata-catalog/incidents/statusWithType"},
            "customFields": {"type": "object"},
        })))
        (extract / "schema-sub").mkdir()
        (extract / "schema-sub" / "incidents.customFields.c.json").write_text(json.dumps({"properties": {
            "free_text": {"type": ["string", "null"], "maxLength": 100},
        }}))
        (extract / "named-ids").mkdir()
        (extract / "named-ids" / "incidents.customFields.c.special_code_1.json").write_text(json.dumps({"items": [{"id": 1, "lookupName": "A"}]}))
        (extract / "named-ids" / "incidents.statusWithType.status.json").write_text(json.dumps({"items": [{"id": 1, "lookupName": "Solved"}]}))
        (extract / "named-ids" / "hierarchy.incidents.source.json").write_text(json.dumps({"items": [{"id": 2, "lookupName": "Phone"}]}))
        (extract / "menus").mkdir()
        (extract / "menus" / "serviceCategories.json").write_text(json.dumps({"items": [{"id": 3, "lookupName": "Cat"}]}))
        rows = {r["field_path"]: r for r in bdd.build(extract).to_dicts()}
        assert rows["AfterMeals"]["is_menu"] is False and rows["AfterMeals"]["suggested_sf_type"] == "Checkbox"
        assert rows["createdByAccount"]["is_menu"] is False and rows["createdByAccount"]["suggested_sf_type"] == "Lookup(User)"
        assert rows["category"]["menu_source"] == "menus/serviceCategories"
        assert rows["source"]["menu_source"] == "named-ids/hierarchy.incidents.source"
        assert rows["statusWithType"]["menu_source"] == "named-ids/incidents.statusWithType.status"
        # a plain custom text field never inherits another custom field's option list
        assert rows["customFields.c.free_text"]["menu_source"] == "" and rows["customFields.c.free_text"]["is_menu"] is False


def test_dates_are_read_from_oracle_bounds():
    r = bdd._row("incidents", "createdTime", "standard", {"type": "string", "minimumDateTime": "1970-01-02T00:00:00.000Z"}, "")
    assert r["suggested_sf_type"] == "Date/Time"
    r = bdd._row("incidents", "customFields.c.last_smoked", "custom (c)", {"type": ["string", "null"], "minimumDate": "1970-01-02"}, "")
    assert r["suggested_sf_type"] == "Date"


def test_subtype_schemas_pulled_and_expanded(monkeypatch):
    calls: list[str] = []
    base = "https://x/services/rest/connect/v1.4/metadata-catalog"
    tasks = _oracle_wrapped({
        "id": {"type": "integer"},
        "serviceSettings": {"type": "object", "$ref": f"{base}/tasks/serviceSettings"},
        "contact": {"type": "object", "$ref": f"{base}/contacts"},  # another object: never followed
    })
    sub = {"properties": {"incident": {"type": ["object", "null"], "$ref": f"{base}/incidents"},
                          "detail": {"type": "object", "$ref": f"{base}/tasks/serviceSettings/detail"}}}
    sub2 = {"properties": {"when": {"type": "string", "minimumDate": "1970-01-02"}}}

    class FakeSession:
        headers: dict = {}

        def get(self, url, auth=None, timeout=None, headers=None, params=None):
            rel = url.split("/services/rest/connect/v1.4/", 1)[1]
            calls.append(rel)
            assert headers == {"Accept": "application/schema+json"}
            return _FakeResp({"metadata-catalog/tasks/serviceSettings": sub, "metadata-catalog/tasks/serviceSettings/detail": sub2}[rel])

    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp) / "2026-09-23"
        (out / "schema").mkdir(parents=True)
        (out / "schema" / "tasks.json").write_text(json.dumps(tasks))
        client = oc.OsvcClient("https://example.invalid", ("u", "p"), out / "_manifest.jsonl", session=FakeSession())
        scope: dict = {}
        pm.pull_subtype_schemas(client, out, {}, scope)
        assert calls == ["metadata-catalog/tasks/serviceSettings", "metadata-catalog/tasks/serviceSettings/detail"]
        assert scope["subtype_schemas_fetched"] == 2
        rows = {r["field_path"]: r for r in bdd.build(out).to_dicts()}
        assert "serviceSettings" not in rows and "serviceSettings.detail" not in rows
        assert rows["serviceSettings.incident"]["lookup_target"] == "incidents"
        assert rows["serviceSettings.incident"]["suggested_sf_type"] == "Lookup"
        assert rows["serviceSettings.detail.when"]["suggested_sf_type"] == "Date"
        assert rows["contact"]["suggested_sf_type"] == "Lookup"


# --- component registry and the new configuration components ---------------------------------

import components as comp  # noqa: E402


def test_registry_order_and_dependencies_are_registered_first():
    names = list(comp.REGISTRY)
    assert names[0] == "catalog"
    for c in comp.REGISTRY.values():
        for dep in c.depends_on:
            assert names.index(dep) < names.index(c.name), (c.name, dep)
    assert {"service-menus", "report-definitions", "mailboxes", "standard-content", "event-subscriptions"} <= set(names)


def test_plan_adds_missing_dependencies_and_skips_satisfied_ones():
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp)
        # empty folder: report-definitions pulls in report-list and catalog
        assert comp.plan(["report-definitions"], out, {}) == ["catalog", "report-list", "report-definitions"]
        # the report list is already there: only the selected component runs
        (out / "menus").mkdir()
        (out / "menus" / "analyticsReports.json").write_text("{}")
        assert comp.plan(["report-definitions"], out, {}) == ["report-definitions"]
        # a component recorded in _scope.json counts as done
        assert comp.plan(["service-menus"], out, {"components_run": {"menu-listings": "x"}}) == ["service-menus"]
        # an explicitly selected component runs even when done; its missing catalog comes with it
        assert comp.plan(["report-list"], out, {}) == ["catalog", "report-list"]
        try:
            comp.plan(["no-such-thing"], out, {})
        except ValueError as exc:
            assert "Unknown components" in str(exc)
        else:
            raise AssertionError("unknown component accepted")


def test_report_targeting_by_keyword_and_date_or_by_ids():
    listing = {"items": [
        {"id": 1, "lookupName": "SMS Report #1", "updatedTime": "2020-07-22T00:00:00Z"},
        {"id": 2, "lookupName": "VA Callbacks", "updatedTime": "2020-09-12T00:00:00Z"},
        {"id": 3, "lookupName": "Incidents by Queue", "updatedTime": "2007-01-01T00:00:00Z"},  # stock, no keyword
        {"id": 4, "lookupName": "Survey results", "updatedTime": "2007-01-01T00:00:00Z"},  # keyword, but stock era
        {"id": 5, "lookupName": "Weekly totals", "updatedTime": "2025-01-01T00:00:00Z"},  # recent, no keyword
    ]}
    assert comp.target_reports(listing) == [1, 2]
    assert comp.target_reports(listing, [3, 5]) == [3, 5]


def _config_server(calls: list):
    base = "https://x/services/rest/connect/v1.4"
    rows = {
        "serviceDispositions": {"items": [{"id": 9}, {"id": 11}]},
        "serviceDispositions/9": {"id": 9, "lookupName": "Break Off", "parent": None,
                                  "adminVisibleInterfaces": {"links": [{"rel": "self", "href": f"{base}/serviceDispositions/9/adminVisibleInterfaces"}]}},
        "serviceDispositions/11": {"id": 11, "lookupName": "Time Constraint", "parent": {"links": [{"rel": "self", "href": f"{base}/serviceDispositions/9"}]},
                                   "adminVisibleInterfaces": {"links": [{"rel": "self", "href": f"{base}/serviceDispositions/11/adminVisibleInterfaces"}]}},
        "serviceDispositions/9/adminVisibleInterfaces": {"items": [{"rel": "canonical", "href": f"{base}/siteInterfaces/1"}, {"rel": "canonical", "href": f"{base}/siteInterfaces/5"}]},
        "serviceDispositions/11/adminVisibleInterfaces": {"items": [{"rel": "canonical", "href": f"{base}/siteInterfaces/1"}]},
        "mailboxes": {"items": [{"id": 1}]},
        "mailboxes/1": {"id": 1, "name": "NCI", "outgoingEmailSettings": {"links": [{"rel": "self", "href": f"{base}/mailboxes/1/outgoingEmailSettings"}]}},
        "mailboxes/1/outgoingEmailSettings": {"fromAddress": "service@example.invalid", "isEnabled": True},
        "standardContents": {"items": [{"id": 52}]},
        "standardContents/52": {"id": 52, "name": "Greeting", "contentValues": {"links": [{"rel": "self", "href": f"{base}/standardContents/52/contentValues"}]}},
        "standardContents/52/contentValues": {"items": [{"rel": "canonical", "href": f"{base}/standardContents/52/contentValues/1"}]},
        "standardContents/52/contentValues/1": {"contentType": {"lookupName": "Text"}, "value": "Thank you for calling."},
        "eventSubscriptions": {"items": [{"id": 7}]},
        "eventSubscriptions/7": {"id": 7, "name": "Push", "endPoint": "https://hook.example.invalid", "integrationUser": {"id": 99, "lookupName": "api_user"}},
        "analyticsReports/101125": {"id": 101125, "lookupName": "SMS Report #1", "createdTime": "2020", "updatedTime": "2020"},
        "analyticsReports/101125/columns": {"items": [{"rel": "canonical", "href": f"{base}/analyticsReports/101125/columns/0"}]},
        "analyticsReports/101125/columns/0": {"heading": "Task ID", "dataType": {"lookupName": "INT"}},
        "analyticsReports/101125/filters": {"items": [{"rel": "canonical", "href": f"{base}/analyticsReports/101125/filters/0"}]},
        "analyticsReports/101125/filters/0": {"name": "SMS Consent", "values": ["1"]},
    }

    class FakeSession:
        headers: dict = {}

        def get(self, url, auth=None, timeout=None, headers=None, params=None):
            rel = url.split("/services/rest/connect/v1.4/", 1)[1]
            calls.append(rel)
            if rel not in rows:
                raise AssertionError(f"unexpected request {rel}")
            return _FakeResp(rows[rel])

    return FakeSession()


def test_configuration_components_against_fake_server():
    calls: list[str] = []
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp)
        client = oc.OsvcClient("https://example.invalid", ("u", "p"), out / "_manifest.jsonl", session=_config_server(calls))
        (out / "catalog.json").write_text(json.dumps({"items": [{"name": n} for n in ("serviceDispositions", "mailboxes", "standardContents", "eventSubscriptions")]}))
        ctx = comp.Context(client, out, {}, {}, report_ids=[101125])

        rows = comp._fetch_rows(ctx, "serviceDispositions", lambda k, v: k in comp.SERVICE_MENU_EXPAND)
        assert rows[1]["parent"]["lookupName"] == "Break Off"
        assert len(rows[0]["adminVisibleInterfaces"]["items"]) == 2

        saved = comp._fetch_rows(ctx, "mailboxes", lambda k, v: True)
        assert saved[0]["outgoingEmailSettings"]["fromAddress"] == "service@example.invalid"

        comp.pull_standard_content(ctx)
        sc = json.loads((out / "config-rows" / "standardContents.json").read_text())
        assert sc[0]["contentValues"][0]["value"] == "Thank you for calling."

        comp.pull_event_subscriptions(ctx)
        ev = json.loads((out / "config-rows" / "eventSubscriptions.json").read_text())
        assert "integrationUser" not in ev[0] and ev[0]["endPoint"] == "https://hook.example.invalid"

        comp.pull_report_definitions(ctx)
        rd = json.loads((out / "report-defs" / "101125.full.json").read_text())
        assert rd["columns"][0]["heading"] == "Task ID" and rd["filters"][0]["name"] == "SMS Consent"
        before = len(calls)
        comp.pull_report_definitions(ctx)  # resume: already saved, nothing requested
        assert len(calls) == before and ctx.scope["report_definitions"]["already_present"] == 1

        for rel in calls:
            assert rel.split("/", 1)[0] not in oc.REFUSED_PREFIXES, rel


def test_hierarchy_parents_fetched_once_per_value_and_built_into_paths():
    calls: list[str] = []
    parents = {
        "namedIDHierarchies/incidents/source/1/parents": {"items": []},
        "namedIDHierarchies/incidents/source/2/parents": {"items": [{"id": 1, "lookupName": "Root"}]},
        "namedIDHierarchies/incidents/source/3/parents": {"items": [{"id": 1, "lookupName": "Root"}, {"id": 2, "lookupName": "Mid"}]},
    }

    class FakeSession:
        headers: dict = {}

        def get(self, url, auth=None, timeout=None, headers=None, params=None):
            rel = url.split("/services/rest/connect/v1.4/", 1)[1]
            calls.append(rel)
            _, _res, rest = rel.split("/", 2)  # any object carrying the list answers the same
            return _FakeResp(parents[f"namedIDHierarchies/incidents/{rest}"])

    link = lambda res, i: {"links": [{"rel": "self", "href": f"https://x/services/rest/connect/v1.4/namedIDHierarchies/{res}/source/{i}/parents"}]}
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp)
        (out / "named-ids").mkdir()
        for res in ("incidents", "contacts"):  # the same list on two objects
            (out / "named-ids" / f"hierarchy.{res}.source.json").write_text(json.dumps({"items": [
                {"id": 1, "lookupName": "Root", "parents": link(res, 1)},
                {"id": 2, "lookupName": "Mid", "parents": link(res, 2)},
                {"id": 3, "lookupName": "Leaf", "parents": link(res, 3)},
            ]}))
        client = oc.OsvcClient("https://example.invalid", ("u", "p"), out / "_manifest.jsonl", session=FakeSession())
        ctx = comp.Context(client, out, {}, {})
        comp.pull_hierarchy_parents(ctx)
        assert len(calls) == 3  # once per value, reused for the second object
        assert ctx.scope["hierarchy_parents_fetched"] == 3

        pk = bdd.build_picklists(out)
        leaf = pk.filter((pk["lookup_name"] == "Leaf") & (pk["object"] == "contacts")).to_dicts()[0]
        assert leaf["parent"] == "Mid" and leaf["hierarchy_path"] == "Root > Mid > Leaf"
        root = pk.filter((pk["lookup_name"] == "Root") & (pk["object"] == "incidents")).to_dicts()[0]
        assert root["parent"] == "" and root["hierarchy_path"] == "Root"


# --- configuration inventory workbook ---------------------------------------------

import build_config_inventory as bci  # noqa: E402


def test_inventory_workbook_holds_back_people_names_and_standard_content_text():
    import polars as pl
    from openpyxl import load_workbook

    with tempfile.TemporaryDirectory() as tmp:
        extract = Path(tmp) / "2026-09-23"
        for sub in ("config-rows", "named-ids", "menus", "report-defs"):
            (extract / sub).mkdir(parents=True)
        pl.DataFrame({c: ["incidents"] if c == "object" else ["customFields.c.ct_searcher"] if c == "field_path"
                      else [None] for c in bdd.COLUMNS}, schema_overrides={"is_menu": pl.Boolean, "max_length": pl.Int64,
                      "read_only": pl.Boolean}).write_parquet(extract / "data-dictionary.parquet")
        pl.DataFrame([
            {"source": "customMenu", "object": "incidents", "field_or_object": "customFields.c.ct_searcher", "id": 1, "lookup_name": "Searcherperson"},
            {"source": "customMenu", "object": "incidents", "field_or_object": "customFields.c.lead_used", "id": 2, "lookup_name": "Leadperson"},
            {"source": "customMenu", "object": "incidents", "field_or_object": "customFields.c.service_number", "id": 3, "lookup_name": "4CANCER"},
        ]).with_columns(
            *[pl.lit(None, dtype=pl.String).alias(c) for c in ("parent", "hierarchy_path", "visible_interfaces")],
            pl.lit(None, dtype=pl.Int64).alias("display_order"),
        ).select(bdd.PICKLIST_COLUMNS).write_parquet(extract / "picklist-values.parquet")
        (extract / "menus" / "siteInterfaces.json").write_text(json.dumps({"items": [{"id": 1, "lookupName": "nci"}]}))
        (extract / "named-ids" / "incidents.queue.json").write_text(json.dumps({"items": [
            {"id": 8, "lookupName": "English Cancer Phone Call"}, {"id": 28, "lookupName": "------"}, {"id": 14, "lookupName": "English Cancer Chat"}]}))
        (extract / "config-rows" / "standardContents.json").write_text(json.dumps([{
            "id": 1, "lookupName": "Greeting", "hotKey": "hi", "folder": {"lookupName": "LH SRL", "parents": []},
            "contentValues": [{"contentType": {"id": 1, "lookupName": "Text"}, "value": "Private reply text"}],
            "usage": {"incidentText": True}, "adminVisibleInterfaces": _iface_links("standardContents", 1, [1])}]))
        (extract / "menus" / "analyticsReports.json").write_text(json.dumps({"items": [
            {"id": 100047, "lookupName": "Call Backs", "createdTime": "2012-03-20T00:00:00.000Z", "updatedTime": "2012-03-20T00:00:00.000Z"},
            {"id": 12, "lookupName": "Stock report", "createdTime": "2007-01-01T00:00:00.000Z", "updatedTime": "2007-01-01T00:00:00.000Z"}]}))
        (extract / "report-defs" / "100047.full.json").write_text(json.dumps(
            {"id": 100047, "updatedTime": "2026-09-23T21:11:58.000Z", "columns": [{}, {}], "filters": [{}]}))

        out = bci.write_workbook(extract, bci.build(extract))

        wb = load_workbook(out, read_only=True)
        cells = {c for ws in wb.worksheets for row in ws.iter_rows(values_only=True) for c in row if isinstance(c, str)}
        assert "Searcherperson" not in cells and "Leadperson" not in cells and "4CANCER" in cells
        assert "Private reply text" not in cells
        header = next(wb["Standard content index"].iter_rows(values_only=True))
        assert header == ("id", "name", "folder_path", "hot_key", "content_types", "has_html", "used_in_inquiry_text",
                          "used_in_chat_text", "used_as_chat_url", "used_in_rule_text", "visible_interfaces")
        assert next(wb["Standard content index"].iter_rows(min_row=2, values_only=True))[-1] == "nci"

        queues = bci.queues(extract)
        assert queues["is_separator"].to_list() == [False, True, False] and queues["section"].to_list() == [1, None, 2]
        reports = bci.reports(extract).sort("id")
        assert reports["kind"].to_list() == ["Oracle stock", "custom"]
        assert reports["definition_date_changed_by_pull"].to_list() == [False, True]
        assert reports["updated"].to_list() == ["2007-01-01", "2012-03-20"]


def test_config_named_ids_skip_staff_references_and_save_where_the_build_reads(monkeypatch):
    calls: list[str] = []

    class FakeSession:
        headers: dict = {}

        def get(self, url, auth=None, timeout=None, headers=None, params=None):
            rel = url.split("/services/rest/connect/v1.4/", 1)[1]
            calls.append(rel)
            assert not rel.endswith("integrationUser"), "staff reference was fetched"
            return _FakeResp({"items": [{"id": 1, "lookupName": "Integer"}], "hasMore": False})

    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp)
        client = oc.OsvcClient("https://example.invalid", ("u", "p"), out / "_manifest.jsonl", session=FakeSession())
        ctx = comp.Context(client=client, out_dir=out, errors={}, scope={})
        comp.REGISTRY["config-named-ids"].run(ctx)
        assert len(calls) == len(comp.CONFIG_NAMED_ID_LISTS)
        assert all(c.startswith("namedIDs/") for c in calls)
        assert (out / "named-ids" / "configurations.dataType.json").exists()
        assert (out / "named-ids" / "organizations.salesSettings.totalRevenue.currency.json").exists()
        assert ctx.scope["config_named_id_lists_fetched"] == len(comp.CONFIG_NAMED_ID_LISTS)
        rows = bdd.build_picklists(out)
        assert ("configurations", "dataType") in set(zip(rows["object"], rows["field_or_object"]))
